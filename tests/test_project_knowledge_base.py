import io
from unittest.mock import MagicMock, patch

import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

import project_knowledge_base as pkb


def _docx_bytes(paragraphs):
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pptx_bytes(slide_texts):
    prs = Presentation()
    layout = prs.slide_layouts[6]
    for text in slide_texts:
        slide = prs.slides.add_slide(layout)
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        box.text_frame.text = text
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_extract_text_from_docx_joins_paragraphs():
    file_bytes = _docx_bytes(["First paragraph.", "Second paragraph."])

    result = pkb.extract_text_from_docx(file_bytes)

    assert result == "First paragraph.\n\nSecond paragraph."


def test_extract_text_from_pptx_joins_slide_text():
    file_bytes = _pptx_bytes(["Slide one text", "Slide two text"])

    result = pkb.extract_text_from_pptx(file_bytes)

    assert "Slide one text" in result
    assert "Slide two text" in result


def test_extract_text_from_txt_decodes_utf8():
    result = pkb.extract_text_from_txt("Hello world".encode("utf-8"))
    assert result == "Hello world"


@patch("project_knowledge_base.PdfReader")
def test_extract_text_from_pdf_joins_page_text(mock_reader_cls):
    page1, page2 = MagicMock(), MagicMock()
    page1.extract_text.return_value = "Page one text"
    page2.extract_text.return_value = "Page two text"
    mock_reader_cls.return_value.pages = [page1, page2]

    result = pkb.extract_text_from_pdf(b"fake-pdf-bytes")

    assert result == "Page one text\n\nPage two text"


def test_extract_text_dispatches_on_source_format():
    result = pkb.extract_text("Hello world".encode("utf-8"), "txt")
    assert result == "Hello world"


def test_extract_text_raises_on_unsupported_format():
    with pytest.raises(ValueError, match="audio"):
        pkb.extract_text(b"data", "audio")


def test_chunk_text_splits_on_paragraph_breaks():
    text = "First paragraph.\n\nSecond paragraph.\n\nThird paragraph."

    result = pkb.chunk_text(text)

    assert result == ["First paragraph.", "Second paragraph.", "Third paragraph."]


def test_chunk_text_splits_long_paragraph_into_fixed_size_pieces():
    long_paragraph = "x" * 2500

    result = pkb.chunk_text(long_paragraph, max_chunk_chars=1000)

    assert result == ["x" * 1000, "x" * 1000, "x" * 500]


def test_infer_source_format_maps_known_document_extensions():
    assert pkb.infer_source_format("report.pdf") == "pdf"
    assert pkb.infer_source_format("notes.docx") == "docx"
    assert pkb.infer_source_format("deck.pptx") == "pptx"
    assert pkb.infer_source_format("transcript.txt") == "txt"


def test_infer_source_format_maps_markdown_extension():
    assert pkb.infer_source_format("readme.md") == "md"


def test_extract_text_dispatches_markdown_to_txt_extractor():
    result = pkb.extract_text("# Heading\n\nBody text.".encode("utf-8"), "md")
    assert result == "# Heading\n\nBody text."


def test_infer_source_format_maps_audio_extensions():
    assert pkb.infer_source_format("meeting.mp3") == "audio"
    assert pkb.infer_source_format("call.wav") == "audio"


def test_infer_source_format_raises_on_unsupported_extension():
    with pytest.raises(ValueError, match="exe"):
        pkb.infer_source_format("virus.exe")


def test_infer_artifact_type_maps_audio_and_document():
    assert pkb.infer_artifact_type("audio") == "audio"
    assert pkb.infer_artifact_type("pdf") == "document"
    assert pkb.infer_artifact_type("docx") == "document"


from google.auth.exceptions import DefaultCredentialsError


def test_transcribe_audio_skips_when_bucket_not_configured(monkeypatch):
    monkeypatch.delenv("GCS_TRANSCRIBE_BUCKET", raising=False)

    with patch("project_knowledge_base.storage.Client") as mock_storage:
        result = pkb.transcribe_audio(b"fake-audio-bytes", "meeting.mp3")

    assert result is None
    mock_storage.assert_not_called()


def test_transcribe_audio_skips_unsupported_audio_format(monkeypatch):
    monkeypatch.setenv("GCS_TRANSCRIBE_BUCKET", "cosmos-transcribe-bucket")

    with patch("project_knowledge_base.storage.Client") as mock_storage:
        result = pkb.transcribe_audio(b"fake-audio-bytes", "meeting.m4a")

    assert result is None
    mock_storage.assert_not_called()


@patch("project_knowledge_base.storage.Client")
def test_transcribe_audio_returns_none_when_credentials_missing(mock_storage_cls, monkeypatch):
    monkeypatch.setenv("GCS_TRANSCRIBE_BUCKET", "cosmos-transcribe-bucket")
    mock_storage_cls.side_effect = DefaultCredentialsError()

    result = pkb.transcribe_audio(b"fake-audio-bytes", "meeting.mp3")

    assert result is None


def _mock_storage_client():
    mock_storage_client = MagicMock()
    mock_blob = MagicMock()
    mock_storage_client.bucket.return_value.blob.return_value = mock_blob
    return mock_storage_client, mock_blob


@patch("project_knowledge_base.speech.SpeechClient")
@patch("project_knowledge_base.storage.Client")
def test_transcribe_audio_returns_transcript_on_completed_job(mock_storage_cls, mock_speech_cls):
    import os
    os.environ["GCS_TRANSCRIBE_BUCKET"] = "cosmos-transcribe-bucket"
    try:
        mock_storage_client, mock_blob = _mock_storage_client()
        mock_storage_cls.return_value = mock_storage_client

        mock_speech_client = MagicMock()
        mock_speech_cls.return_value = mock_speech_client
        mock_operation = MagicMock()
        mock_operation.done.return_value = True
        mock_result_alt = MagicMock()
        mock_result_alt.transcript = "hello from the meeting"
        mock_result = MagicMock()
        mock_result.alternatives = [mock_result_alt]
        mock_operation.result.return_value.results = [mock_result]
        mock_speech_client.long_running_recognize.return_value = mock_operation

        result = pkb.transcribe_audio(b"fake-audio-bytes", "meeting.mp3")

        assert result == "hello from the meeting"
        mock_blob.upload_from_string.assert_called_once_with(b"fake-audio-bytes")
        mock_speech_client.long_running_recognize.assert_called_once()
    finally:
        del os.environ["GCS_TRANSCRIBE_BUCKET"]


@patch("project_knowledge_base.speech.SpeechClient")
@patch("project_knowledge_base.storage.Client")
def test_transcribe_audio_returns_none_when_job_raises(mock_storage_cls, mock_speech_cls, monkeypatch):
    from google.api_core.exceptions import GoogleAPICallError

    monkeypatch.setenv("GCS_TRANSCRIBE_BUCKET", "cosmos-transcribe-bucket")

    mock_storage_client, mock_blob = _mock_storage_client()
    mock_storage_cls.return_value = mock_storage_client

    mock_speech_client = MagicMock()
    mock_speech_cls.return_value = mock_speech_client
    mock_operation = MagicMock()
    mock_operation.done.return_value = True
    mock_operation.result.side_effect = GoogleAPICallError("job failed")
    mock_speech_client.long_running_recognize.return_value = mock_operation

    result = pkb.transcribe_audio(b"fake-audio-bytes", "meeting.mp3")

    assert result is None


@patch("project_knowledge_base.speech.SpeechClient")
@patch("project_knowledge_base.storage.Client")
def test_transcribe_audio_returns_none_when_no_speech_detected(mock_storage_cls, mock_speech_cls, monkeypatch):
    monkeypatch.setenv("GCS_TRANSCRIBE_BUCKET", "cosmos-transcribe-bucket")

    mock_storage_client, mock_blob = _mock_storage_client()
    mock_storage_cls.return_value = mock_storage_client

    mock_speech_client = MagicMock()
    mock_speech_cls.return_value = mock_speech_client
    mock_operation = MagicMock()
    mock_operation.done.return_value = True
    mock_operation.result.return_value.results = []
    mock_speech_client.long_running_recognize.return_value = mock_operation

    result = pkb.transcribe_audio(b"fake-audio-bytes", "meeting.mp3")

    assert result is None


@patch("project_knowledge_base.time.sleep")
@patch("project_knowledge_base.speech.SpeechClient")
@patch("project_knowledge_base.storage.Client")
def test_transcribe_audio_returns_none_when_polling_window_exceeded(mock_storage_cls, mock_speech_cls, mock_sleep, monkeypatch):
    monkeypatch.setenv("GCS_TRANSCRIBE_BUCKET", "cosmos-transcribe-bucket")

    mock_storage_client, mock_blob = _mock_storage_client()
    mock_storage_cls.return_value = mock_storage_client

    mock_speech_client = MagicMock()
    mock_speech_cls.return_value = mock_speech_client
    mock_operation = MagicMock()
    mock_operation.done.return_value = False
    mock_speech_client.long_running_recognize.return_value = mock_operation

    result = pkb.transcribe_audio(b"fake-audio-bytes", "meeting.mp3")

    assert result is None


_DOCUMENT_ARTIFACT = {
    "id": 1, "project_id": 10, "filename": "notes.txt", "artifact_type": "document",
    "source_format": "txt", "purpose": "reference", "status": "Uploaded",
    "transcript_text": None, "uploaded_by": 5, "uploaded_at": "2026-08-28T09:00:00",
}
_AUDIO_ARTIFACT = {
    "id": 2, "project_id": 10, "filename": "meeting.mp3", "artifact_type": "audio",
    "source_format": "audio", "purpose": "reference", "status": "Uploaded",
    "transcript_text": None, "uploaded_by": 5, "uploaded_at": "2026-08-28T09:00:00",
}


def _fake_conn():
    cursor = MagicMock()
    cursor.__enter__.return_value = cursor
    cursor.__exit__.return_value = False
    conn = MagicMock()
    conn.cursor.return_value = cursor
    return conn, cursor


def _fake_rag():
    rag = MagicMock()
    rag.embedding_model.encode.return_value = [[0.1, 0.2], [0.3, 0.4]]
    return rag


@patch("project_knowledge_base.get_db_connection")
def test_insert_chunks_inserts_one_row_per_chunk(mock_get_conn):
    conn, cursor = _fake_conn()
    mock_get_conn.return_value = conn

    pkb._insert_chunks(10, 1, ["chunk one", "chunk two"], [[0.1, 0.2], [0.3, 0.4]])

    assert cursor.execute.call_count == 2
    first_sql, first_params = cursor.execute.call_args_list[0][0]
    assert "INSERT INTO project_kb_chunks" in first_sql
    assert first_params == (10, 1, "chunk one", [0.1, 0.2])
    conn.commit.assert_called_once()


@patch("project_knowledge_base._insert_chunks")
@patch("project_knowledge_base.chunk_text", return_value=["chunk one", "chunk two"])
@patch("project_knowledge_base.extract_text", return_value="parsed document text")
@patch("project_knowledge_base.project_artifacts_db.get_artifact_by_id", return_value=_DOCUMENT_ARTIFACT)
@patch(
    "project_knowledge_base.project_artifacts_db.update_artifact_status",
    return_value={**_DOCUMENT_ARTIFACT, "status": "Indexed"},
)
def test_ingest_artifact_indexes_a_document(mock_update, mock_get, mock_extract, mock_chunk, mock_insert):
    rag = _fake_rag()

    result = pkb.ingest_artifact(rag, 1, b"file-bytes")

    assert result["status"] == "Indexed"
    mock_extract.assert_called_once_with(b"file-bytes", "txt")
    mock_chunk.assert_called_once_with("parsed document text")
    mock_insert.assert_called_once_with(10, 1, ["chunk one", "chunk two"], [[0.1, 0.2], [0.3, 0.4]])
    mock_update.assert_any_call(1, "Processing")
    mock_update.assert_any_call(1, "Indexed", transcript_text=None)


@patch("project_knowledge_base._insert_chunks")
@patch("project_knowledge_base.chunk_text", return_value=["chunk one"])
@patch("project_knowledge_base.transcribe_audio", return_value="hello from the meeting")
@patch("project_knowledge_base.project_artifacts_db.get_artifact_by_id", return_value=_AUDIO_ARTIFACT)
@patch(
    "project_knowledge_base.project_artifacts_db.update_artifact_status",
    return_value={**_AUDIO_ARTIFACT, "status": "Indexed", "transcript_text": "hello from the meeting"},
)
def test_ingest_artifact_indexes_audio_with_transcript(mock_update, mock_get, mock_transcribe, mock_chunk, mock_insert):
    rag = _fake_rag()

    result = pkb.ingest_artifact(rag, 2, b"audio-bytes")

    assert result["status"] == "Indexed"
    assert result["transcript_text"] == "hello from the meeting"
    mock_transcribe.assert_called_once_with(b"audio-bytes", "meeting.mp3")
    mock_chunk.assert_called_once_with("hello from the meeting")
    mock_update.assert_any_call(2, "Indexed", transcript_text="hello from the meeting")


@patch("project_knowledge_base.transcribe_audio", return_value=None)
@patch("project_knowledge_base.project_artifacts_db.get_artifact_by_id", return_value=_AUDIO_ARTIFACT)
@patch(
    "project_knowledge_base.project_artifacts_db.update_artifact_status",
    return_value={**_AUDIO_ARTIFACT, "status": "Transcript Needed"},
)
def test_ingest_artifact_marks_transcript_needed_when_gcs_unavailable(mock_update, mock_get, mock_transcribe):
    rag = _fake_rag()

    result = pkb.ingest_artifact(rag, 2, b"audio-bytes")

    assert result["status"] == "Transcript Needed"
    mock_transcribe.assert_called_once_with(b"audio-bytes", "meeting.mp3")
    mock_update.assert_any_call(2, "Transcript Needed")


@patch("project_knowledge_base.extract_text", side_effect=RuntimeError("corrupt file"))
@patch("project_knowledge_base.project_artifacts_db.get_artifact_by_id", return_value=_DOCUMENT_ARTIFACT)
@patch(
    "project_knowledge_base.project_artifacts_db.update_artifact_status",
    return_value={**_DOCUMENT_ARTIFACT, "status": "Failed"},
)
def test_ingest_artifact_marks_failed_on_parse_error(mock_update, mock_get, mock_extract):
    rag = _fake_rag()

    result = pkb.ingest_artifact(rag, 1, b"file-bytes")

    assert result["status"] == "Failed"
    mock_update.assert_any_call(1, "Failed")


@patch("project_knowledge_base._insert_chunks")
@patch("project_knowledge_base.chunk_text", return_value=[])
@patch("project_knowledge_base.extract_text", return_value="")
@patch("project_knowledge_base.project_artifacts_db.get_artifact_by_id", return_value=_DOCUMENT_ARTIFACT)
@patch(
    "project_knowledge_base.project_artifacts_db.update_artifact_status",
    return_value={**_DOCUMENT_ARTIFACT, "status": "Failed"},
)
def test_ingest_artifact_marks_failed_when_no_text_extracted(mock_update, mock_get, mock_extract, mock_chunk, mock_insert):
    rag = _fake_rag()

    result = pkb.ingest_artifact(rag, 1, b"file-bytes")

    assert result["status"] == "Failed"
    mock_update.assert_any_call(1, "Failed")
    mock_insert.assert_not_called()


def test_ingest_artifact_raises_value_error_for_unknown_artifact():
    with patch("project_knowledge_base.project_artifacts_db.get_artifact_by_id", return_value=None):
        with pytest.raises(ValueError, match="999"):
            pkb.ingest_artifact(_fake_rag(), 999, b"bytes")
